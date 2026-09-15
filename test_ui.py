import os
os.environ['QT_QPA_PLATFORM'] = 'offscreen'
os.environ['MPLCONFIGDIR'] = os.path.abspath('.mplconfig')
import sys
sys.path.insert(0, '.test-deps')
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from openpyxl import load_workbook
from main import DataApp, QApplication, QFileDialog, QMessageBox

app = QApplication.instance() or QApplication([])
app.setStyle('Fusion')

class DashboardTests(unittest.TestCase):
    def setUp(self):
        self.window = DataApp(':memory:')

    def tearDown(self):
        self.window.close()

    def add(self, category, value):
        self.window.category_input.setText(category)
        self.window.value_input.setText(value)
        self.window.add_data()

    def test_filter_and_delete_by_id(self):
        self.add('Satış', '1250,50')
        self.add('Gider', '-250')
        self.add('Satış', '700')
        w = self.window
        w.category_filter.setCurrentIndex(w.category_filter.findData('Satış'))
        self.assertEqual(w.count_card.value.text(), '2')
        self.assertEqual(w.total_card.value.text(), '1.950,50')
        w.data_table.selectRow(0)
        w.delete_selected()
        self.assertEqual([(r[1], r[2]) for r in w.get_data()], [('Satış', 1250.5), ('Gider', -250)])
        w.category_filter.setCurrentIndex(w.category_filter.findData('Gider'))
        self.assertEqual(w.max_card.value.text(), '-250,00')

    def test_validation_and_empty_charts(self):
        for category, value in [('', '12'), ('Test', 'nan'), ('Test', 'inf'), ('Test', 'abc')]:
            self.add(category, value)
        self.assertEqual(self.window.get_data(), [])
        for value in ['0', '-10']:
            self.add('Test', value)
            for chart in (self.window.pie_chart, self.window.bar_chart, self.window.trend_chart):
                chart.canvas.draw()
        with patch.object(QMessageBox, 'question', return_value=QMessageBox.No):
            self.window.delete_all()
        self.assertEqual(len(self.window.get_data()), 2)
        with patch.object(QMessageBox, 'question', return_value=QMessageBox.Yes):
            self.window.delete_all()
        self.assertEqual(self.window.get_data(), [])

    def test_excel_all_records(self):
        self.add('=1+1', '12')
        self.add('Gider', '-2')
        self.window.category_filter.setCurrentIndex(1)
        with TemporaryDirectory() as directory:
            path = str(Path(directory) / 'export.xlsx')
            with patch.object(QFileDialog, 'getSaveFileName', return_value=(path, '')):
                self.window.export_excel()
            wb = load_workbook(path)
            self.assertEqual(wb.active.max_row, 3)
            self.assertEqual(wb.active['A2'].data_type, 's')
            wb.close()

    def test_render_layout(self):
        w = self.window
        records = [('Satış', 18500, '01/09/2026 10:00:00'),
                   ('Hizmet', 9200, '02/09/2026 10:00:00'),
                   ('Abonelik', 6400, '03/09/2026 10:00:00'),
                   ('Satış', 22400, '04/09/2026 10:00:00'),
                   ('Hizmet', 12800, '05/09/2026 10:00:00'),
                   ('Abonelik', 7800, '06/09/2026 10:00:00')]
        w.cursor.executemany('INSERT INTO data (category, value, time) VALUES (?, ?, ?)', records)
        w.conn.commit()
        w.refresh_all()
        w.resize(1400, 1080)
        w.show()
        app.processEvents()
        for chart in (w.pie_chart, w.bar_chart, w.trend_chart):
            chart.canvas.draw()
        self.assertTrue(w.grab().save('ui-preview.png'))
        w.resize(1000, 720)
        app.processEvents()
        self.assertTrue(w.grab().save('ui-preview-small.png'))

if __name__ == '__main__':
    unittest.main()
