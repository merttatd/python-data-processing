import sys
import math
import sqlite3
import hashlib
from pathlib import Path
from datetime import datetime
from PyQt5.QtWidgets import *
from PyQt5.QtCore import Qt
from PyQt5.QtGui import QFont
from matplotlib.backends.backend_qt5agg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.figure import Figure
from matplotlib import dates as mdates
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

PALETTE = ['#6366f1', '#14b8a6', '#f59e0b', '#ec4899', '#3b82f6', '#8b5cf6', '#84cc16']


def number(value):
    return f'{value:,.2f}'.replace(',', '_').replace('.', ',').replace('_', '.')


def label(text, name='muted'):
    widget = QLabel(text)
    widget.setObjectName(name)
    return widget


class Card(QFrame):
    def __init__(self, title, hint, accent):
        super().__init__()
        self.setObjectName('card')
        layout = QVBoxLayout(self)
        layout.setContentsMargins(22, 18, 22, 18)
        layout.setSpacing(7)
        self.title = label(title)
        self.value = label('0', 'metric')
        self.value.setStyleSheet(f'color: {accent};')
        self.hint = label(hint)
        for widget in (self.title, self.value, self.hint):
            layout.addWidget(widget)


class ChartWidget(QFrame):
    def __init__(self, title, subtitle, height=210):
        super().__init__()
        self.setObjectName('card')
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 16, 20, 8)
        layout.setSpacing(4)
        layout.addWidget(label(title, 'sectionTitle'))
        layout.addWidget(label(subtitle))
        self.figure = Figure(facecolor='white', constrained_layout=True)
        self.canvas = FigureCanvas(self.figure)
        self.canvas.setMinimumHeight(height)
        layout.addWidget(self.canvas)

    def axes(self, empty=None):
        self.figure.clear()
        ax = self.figure.add_subplot(111)
        for spine in ax.spines.values():
            spine.set_visible(False)
        ax.tick_params(colors='#8290a5', labelsize=8, length=0, pad=8)
        if empty:
            ax.set_axis_off()
            ax.text(.5, .5, empty, ha='center', va='center', color='#8290a5',
                    fontsize=9, transform=ax.transAxes)
        return ax


class DataApp(QWidget):
    def __init__(self, db_path=None):
        super().__init__()
        self.setWindowTitle('Data Process | Veri Paneli')
        self.resize(1400, 940)
        self.setMinimumSize(1000, 720)
        self.conn = sqlite3.connect(str(db_path or Path(__file__).with_name('data.db')))
        self.cursor = self.conn.cursor()
        self.create_table()
        self.init_ui()
        self.refresh_all()

    def create_table(self):
        self.cursor.execute('''CREATE TABLE IF NOT EXISTS data (
            id INTEGER PRIMARY KEY AUTOINCREMENT, category TEXT, value REAL, time TEXT)''')
        self.conn.commit()

    def init_ui(self):
        self.setFont(QFont('Segoe UI', 10))
        root = QHBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)
        sidebar = QFrame()
        sidebar.setObjectName('sidebar')
        sidebar.setFixedWidth(250)
        side = QVBoxLayout(sidebar)
        side.setContentsMargins(24, 30, 24, 24)
        side.setSpacing(12)
        side.addWidget(label('◈  DATA PROCESS', 'brand'))
        side.addWidget(label('Verilerinize net bir bakış.', 'sideMuted'))
        side.addSpacing(25)
        side.addWidget(label('  ▦   Genel bakış', 'activeNav'))
        side.addSpacing(24)
        side.addWidget(label('YENİ KAYIT', 'sideCaption'))
        side.addWidget(label('Kategori', 'sideLabel'))
        self.category_input = QLineEdit()
        self.category_input.setPlaceholderText('Örn. Satış')
        self.category_input.setMaxLength(120)
        side.addWidget(self.category_input)
        side.addWidget(label('Değer', 'sideLabel'))
        self.value_input = QLineEdit()
        self.value_input.setPlaceholderText('Örn. 1250,50')
        self.value_input.returnPressed.connect(self.add_data)
        side.addWidget(self.value_input)
        add_btn = QPushButton('+  Kayıt ekle')
        add_btn.setObjectName('primary')
        add_btn.clicked.connect(self.add_data)
        side.addWidget(add_btn)
        self.form_message = label('', 'formMessage')
        self.form_message.setWordWrap(True)
        side.addWidget(self.form_message)
        side.addStretch()
        side.addWidget(label('●  Yerel çalışma alanı', 'sideLabel'))
        side.addWidget(label('Veriler bu bilgisayardaki\nSQLite veritabanında saklanır.', 'sideMuted'))
        root.addWidget(sidebar)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        content = QWidget()
        content.setObjectName('workspace')
        main = QVBoxLayout(content)
        main.setContentsMargins(28, 26, 28, 22)
        main.setSpacing(18)
        header = QHBoxLayout()
        headings = QVBoxLayout()
        headings.setSpacing(5)
        headings.addWidget(label('ÇALIŞMA ALANI / ANALİTİK', 'eyebrow'))
        headings.addWidget(label('Genel bakış', 'pageTitle'))
        headings.addWidget(label('Kayıtlarınızı izleyin, karşılaştırın ve yönetin.'))
        header.addLayout(headings)
        header.addStretch()
        export = QPushButton('↓  Excel’e aktar')
        export.setToolTip('Tüm kayıtları Excel dosyası olarak kaydet')
        export.clicked.connect(self.export_excel)
        header.addWidget(export)
        main.addLayout(header)
        filters = QHBoxLayout()
        filters.addWidget(label('Kategori', 'sectionTitle'))
        self.category_filter = QComboBox()
        self.category_filter.setMinimumWidth(180)
        self.category_filter.currentIndexChanged.connect(self.refresh_all)
        filters.addWidget(self.category_filter)
        filters.addStretch()
        self.scope_label = label('')
        filters.addWidget(self.scope_label)
        main.addLayout(filters)
        cards = QHBoxLayout()
        cards.setSpacing(16)
        self.total_card = Card('TOPLAM DEĞER', 'Seçili kapsamdaki toplam', '#6366f1')
        self.count_card = Card('KAYIT SAYISI', 'Seçili kapsamdaki kayıtlar', '#0f766e')
        self.max_card = Card('EN YÜKSEK DEĞER', 'Henüz kayıt yok', '#d97706')
        for card in (self.total_card, self.count_card, self.max_card):
            cards.addWidget(card, 1)
        main.addLayout(cards)
        charts = QGridLayout()
        charts.setSpacing(16)
        self.pie_chart = ChartWidget('Kategori dağılımı', 'Kategori toplamlarının payı')
        self.bar_chart = ChartWidget('Kategori karşılaştırması', 'Toplam değere göre ilk 8 kategori')
        self.trend_chart = ChartWidget('Zaman içindeki değişim', 'Kayıt değerlerinin kronolojik görünümü', 155)
        charts.addWidget(self.pie_chart, 0, 0)
        charts.addWidget(self.bar_chart, 0, 1)
        charts.addWidget(self.trend_chart, 1, 0, 1, 2)
        charts.setColumnStretch(0, 1)
        charts.setColumnStretch(1, 1)
        main.addLayout(charts)
        records = QFrame()
        records.setObjectName('card')
        record_layout = QVBoxLayout(records)
        record_layout.setContentsMargins(20, 16, 20, 16)
        toolbar = QHBoxLayout()
        toolbar.addWidget(label('Kayıtlar', 'sectionTitle'))
        toolbar.addStretch()
        self.delete_btn = QPushButton('Seçileni sil')
        self.delete_btn.clicked.connect(self.delete_selected)
        toolbar.addWidget(self.delete_btn)
        delete_all = QPushButton('Tümünü sil')
        delete_all.setObjectName('danger')
        delete_all.clicked.connect(self.delete_all)
        toolbar.addWidget(delete_all)
        record_layout.addLayout(toolbar)
        self.data_table = QTableWidget(0, 3)
        self.data_table.setHorizontalHeaderLabels(['KATEGORİ', 'DEĞER', 'KAYIT TARİHİ'])
        self.data_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.data_table.verticalHeader().hide()
        self.data_table.verticalHeader().setDefaultSectionSize(42)
        self.data_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.data_table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.data_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.data_table.setShowGrid(False)
        self.data_table.setAlternatingRowColors(True)
        self.data_table.setMinimumHeight(190)
        self.data_table.itemSelectionChanged.connect(
            lambda: self.delete_btn.setEnabled(bool(self.data_table.selectedItems())))
        record_layout.addWidget(self.data_table)
        self.empty_label = label('Henüz kayıt yok. Sol panelden ilk kaydınızı ekleyin.')
        record_layout.addWidget(self.empty_label)
        main.addWidget(records)
        self.status_label = label('')
        main.addWidget(self.status_label)
        scroll.setWidget(content)
        root.addWidget(scroll, 1)
        self.setStyleSheet('''
            QWidget { color: #243149; font-family: 'Segoe UI'; font-size: 13px; }
            QWidget#workspace, QScrollArea { background: #f4f6fb; }
            QFrame#sidebar { background: #151e32; }
            QLabel { background: transparent; }
            #brand { color: white; font-size: 19px; font-weight: 700; }
            #sideMuted { color: #97a5bd; font-size: 12px; }
            #sideCaption { color: #8b9bb6; font-size: 11px; font-weight: 600; }
            #sideLabel { color: #cbd5e1; }
            #activeNav { background: #293356; color: #c9caff; padding: 14px 8px; border-radius: 8px; font-weight: 600; }
            #formMessage { color: #c7d2fe; font-size: 12px; }
            #pageTitle { font-size: 30px; font-weight: 700; color: #17243b; }
            #eyebrow { color: #8793a8; font-size: 10px; font-weight: 600; }
            #muted { color: #7b879c; font-size: 12px; }
            #sectionTitle { font-size: 14px; font-weight: 600; }
            #metric { font-size: 28px; font-weight: 700; }
            QFrame#card { background: white; border: 1px solid #e5e9f2; border-radius: 12px; }
            QPushButton { background: white; border: 1px solid #dfe4ef; border-radius: 7px; padding: 10px 15px; font-weight: 600; }
            QPushButton:hover { background: #eef0ff; border-color: #b6b9f8; }
            QPushButton:pressed { background: #e1e4ff; }
            QPushButton:disabled { color: #aab3c3; background: #f6f7fa; }
            QPushButton#primary { background: #6366f1; color: white; border: 1px solid #6366f1; }
            QPushButton#primary:hover { background: #7477ff; }
            QPushButton#danger { color: #dc5265; border-color: #f0dbe0; }
            QPushButton#danger:hover { background: #fff0f2; }
            QLineEdit, QComboBox { background: white; border: 1px solid #dfe4ef; border-radius: 7px; padding: 10px; }
            QLineEdit:focus, QComboBox:focus, QPushButton:focus { border: 1px solid #818cf8; }
            #sidebar QLineEdit { background: #202c43; color: white; border-color: #35405a; padding: 12px; }
            #sidebar QLineEdit:focus { border-color: #818cf8; }
            QComboBox::drop-down { border: none; width: 24px; }
            QComboBox QAbstractItemView { background: white; selection-background-color: #e9eaff; selection-color: #3730a3; }
            QTableWidget { background: white; alternate-background-color: #f8f9fc; border: none; selection-background-color: #edefff; selection-color: #3730a3; }
            QTableWidget::item { padding: 8px; border-bottom: 1px solid #f0f2f7; }
            QHeaderView::section { background: #f8f9fc; color: #7b879c; border: none; padding: 12px 8px; font-size: 11px; font-weight: 600; }
            QScrollBar:vertical { background: #f4f6fb; width: 10px; margin: 0; }
            QScrollBar::handle:vertical { background: #cbd3e1; border-radius: 5px; min-height: 30px; }
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical { height: 0; }
        ''')

    def get_color(self, category):
        return PALETTE[int(hashlib.sha256(str(category).encode()).hexdigest()[:8], 16) % len(PALETTE)]

    def add_data(self):
        category = self.category_input.text().strip()
        try:
            value = float(self.value_input.text().strip().replace(',', '.'))
            if not category or not math.isfinite(value):
                raise ValueError
        except ValueError:
            self.form_message.setText('Kategori ve geçerli bir sayısal değer girin.')
            return
        self.cursor.execute('INSERT INTO data (category, value, time) VALUES (?, ?, ?)',
                            (category, value, datetime.now().strftime('%d/%m/%Y %H:%M:%S')))
        self.conn.commit()
        self.category_filter.blockSignals(True)
        self.category_filter.setCurrentIndex(0)
        self.category_filter.blockSignals(False)
        self.value_input.clear()
        self.form_message.setText(f'“{category}” kaydı eklendi.')
        self.value_input.setFocus()
        self.refresh_all()

    def get_data(self):
        return self.cursor.execute('SELECT * FROM data ORDER BY id').fetchall()

    def refresh_all(self):
        rows = self.get_data()
        current = self.category_filter.currentData()
        self.category_filter.blockSignals(True)
        self.category_filter.clear()
        self.category_filter.addItem('Tüm kategoriler', None)
        for category in sorted({r[1] for r in rows}, key=lambda c: str(c).casefold()):
            self.category_filter.addItem(str(category), category)
        index = self.category_filter.findData(current)
        self.category_filter.setCurrentIndex(max(0, index))
        self.category_filter.blockSignals(False)
        selected = self.category_filter.currentData()
        self.visible_rows = [r for r in rows if selected is None or r[1] == selected]
        self.data_table.setRowCount(len(self.visible_rows))
        for i, record in enumerate(reversed(self.visible_rows)):
            for j, text in enumerate((str(record[1]), number(record[2]), record[3])):
                item = QTableWidgetItem(text)
                item.setData(Qt.UserRole, record[0])
                if j == 1:
                    item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self.data_table.setItem(i, j, item)
        self.data_table.clearSelection()
        self.delete_btn.setEnabled(False)
        self.empty_label.setVisible(not self.visible_rows)
        self.total_card.value.setText(number(sum(r[2] for r in self.visible_rows)))
        self.count_card.value.setText(str(len(self.visible_rows)))
        maximum = max(self.visible_rows, key=lambda r: r[2], default=None)
        self.max_card.value.setText(number(maximum[2]) if maximum else '—')
        self.max_card.hint.setText(str(maximum[1]) if maximum else 'Henüz kayıt yok')
        self.scope_label.setText(f'{len(self.visible_rows)} / {len(rows)} kayıt gösteriliyor')
        self.status_label.setText(f'●  Güncel  ·  Son yenileme {datetime.now():%H:%M}  ·  Excel aktarımı tüm kayıtları içerir')
        self.draw_all_charts()

    def delete_selected(self):
        index = self.data_table.currentRow()
        if index < 0 or not self.data_table.selectedItems():
            return
        record_id = self.data_table.item(index, 0).data(Qt.UserRole)
        self.cursor.execute('DELETE FROM data WHERE id=?', (record_id,))
        self.conn.commit()
        self.refresh_all()

    def delete_all(self):
        if not self.get_data():
            return
        confirm = QMessageBox.question(self, 'Tüm kayıtları sil',
            'Filtre dışındakiler dahil tüm kayıtlar kalıcı olarak silinecek. Devam edilsin mi?',
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
        if confirm == QMessageBox.Yes:
            self.cursor.execute('DELETE FROM data')
            self.conn.commit()
            self.refresh_all()

    def export_excel(self):
        filename, _ = QFileDialog.getSaveFileName(self, 'Excel dosyasını kaydet',
            f'export_{datetime.now():%Y%m%d_%H%M%S}.xlsx', 'Excel dosyası (*.xlsx)')
        if not filename:
            return
        if not filename.lower().endswith('.xlsx'):
            filename += '.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = 'Veriler'
        ws.append(['Kategori', 'Değer', 'Tarih'])
        for record in self.get_data():
            ws.append([record[1], record[2], record[3]])
            ws.cell(ws.max_row, 1).data_type = 's'
            ws.cell(ws.max_row, 2).number_format = '#,##0.00'
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='6366F1')
        for column, width in [('A', 30), ('B', 20), ('C', 25)]:
            ws.column_dimensions[column].width = width
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        try:
            wb.save(filename)
        except OSError as error:
            QMessageBox.warning(self, 'Dosya kaydedilemedi', str(error))
            return
        self.status_label.setText(f'Excel dosyası kaydedildi: {filename}')

    def aggregate(self):
        data = {}
        for record in self.visible_rows:
            data[record[1]] = data.get(record[1], 0) + record[2]
        return data

    def draw_all_charts(self):
        data = self.aggregate()
        ranked = sorted(data.items(), key=lambda pair: pair[1], reverse=True)
        message = 'Görselleştirmek için bir kayıt ekleyin.'
        pie_ok = bool(data) and min(data.values()) >= 0 and sum(data.values()) > 0
        ax = self.pie_chart.axes(None if pie_ok else (message if not data else
            'Pay dağılımı için toplamlar pozitif olmalı.\nSıfır ve negatif değerler diğer grafiklerde gösterilir.'))
        if pie_ok:
            parts = [(str(k), v) for k, v in ranked if v > 0]
            if len(parts) > 5:
                parts = parts[:5] + [('Diğer kategoriler', sum(v for _, v in parts[5:]))]
            wedges, _ = ax.pie([v for _, v in parts], colors=[self.get_color(k) for k, _ in parts],
                              startangle=90, wedgeprops={'width': .25, 'edgecolor': 'white', 'linewidth': 3})
            ax.text(0, .06, str(len(data)), ha='center', fontsize=25, fontweight='bold', color='#243149')
            ax.text(0, -.2, 'kategori', ha='center', fontsize=9, color='#8290a5')
            total = sum(data.values())
            ax.legend(wedges, [f'{k[:22]}  {v/total:.0%}' for k, v in parts],
                      loc='center left', bbox_to_anchor=(1, .5), frameon=False, fontsize=8, labelcolor='#64748b')
        ax = self.bar_chart.axes(None if data else message)
        if data:
            top = ranked[:8][::-1]
            ax.barh(range(len(top)), [v for _, v in top], height=.5, color=[self.get_color(k) for k, _ in top])
            ax.set_yticks(range(len(top)), [str(k)[:22] for k, _ in top])
            ax.grid(axis='x', color='#edf0f6')
            ax.set_axisbelow(True)
            ax.axvline(0, color='#dfe4ef', linewidth=.8)
            ax.margins(x=.12, y=.2)
        series = {}
        for record in self.visible_rows:
            try:
                date = datetime.strptime(record[3], '%d/%m/%Y %H:%M:%S')
            except (TypeError, ValueError):
                continue
            series.setdefault(record[1], []).append((date, record[2]))
        ax = self.trend_chart.axes(None if series else message)
        if series:
            for category, points in series.items():
                points.sort(key=lambda p: p[0])
                ax.plot([p[0] for p in points], [p[1] for p in points], marker='o', markersize=4,
                        linewidth=2, color=self.get_color(category), label=str(category)[:22])
            locator = mdates.AutoDateLocator(minticks=3, maxticks=7)
            ax.xaxis.set_major_locator(locator)
            ax.xaxis.set_major_formatter(mdates.ConciseDateFormatter(locator))
            ax.grid(axis='y', color='#edf0f6')
            ax.margins(x=.04, y=.25)
            if len(series) <= 5:
                ax.legend(loc='upper left', frameon=False, fontsize=8, ncol=min(5, len(series)))
        for chart in (self.pie_chart, self.bar_chart, self.trend_chart):
            chart.canvas.draw_idle()

    def closeEvent(self, event):
        self.conn.close()
        super().closeEvent(event)


if __name__ == '__main__':
    QApplication.setAttribute(Qt.AA_EnableHighDpiScaling, True)
    QApplication.setAttribute(Qt.AA_UseHighDpiPixmaps, True)
    app = QApplication(sys.argv)
    app.setStyle('Fusion')
    window = DataApp()
    window.show()
    sys.exit(app.exec_())
